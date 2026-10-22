"""Excel export of applied positions (application tracking + VA billing)."""

from __future__ import annotations

import io
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HEADERS = [
    "Applied date", "Applied time", "Applied by", "Company", "Title",
    "Source", "Status", "URL", "Notes", "First seen",
]
COL_WIDTHS = [12, 11, 11, 22, 48, 16, 13, 64, 32, 12]

# Hosts where the company slug is the first path segment.
_PATH_SLUG_HOSTS = ("greenhouse.io", "lever.co", "ashbyhq.com", "workable.com")
# Hosts where the company is the first subdomain label.
_SUBDOMAIN_HOSTS = ("myworkdayjobs.com", "breezy.hr", "recruitee.com", "teamtailor.com")
_GENERIC_LABELS = ("www", "careers", "career", "jobs", "job", "apply", "boards", "app")


def company_from_url(url: str) -> str:
    """Best-effort company name from an ATS/careers-page URL."""
    try:
        p = urlparse(url)
    except ValueError:
        return ""
    host = (p.hostname or "").lower()
    path_parts = [s for s in p.path.split("/") if s]

    if any(h in host for h in _PATH_SLUG_HOSTS):
        slug = path_parts[0] if path_parts else ""
        return slug.replace("-", " ").replace("_", " ").title()
    if any(h in host for h in _SUBDOMAIN_HOSTS):
        return host.split(".")[0].title()

    # Careers pages etc.: strip generic subdomains, use the registrable label.
    labels = host.split(".")
    while len(labels) > 2 and labels[0] in _GENERIC_LABELS:
        labels = labels[1:]
    return labels[0].title() if labels else host


def applied_workbook_bytes(jobs: list[dict]) -> bytes:
    """Build an .xlsx of applied positions, newest application first."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Applied"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    applied = sorted(
        [j for j in jobs if j.get("applied_at")],
        key=lambda j: j["applied_at"],
        reverse=True,
    )

    for j in applied:
        at = j.get("applied_at") or ""
        date, _, time = at.partition("T")
        who = j.get("applied_by") or "admin"  # pre-VA rows were applied by admin
        ws.append([
            date,
            time,
            {"va": "VA", "auto": "Auto"}.get(who, "Admin"),
            company_from_url(j.get("url") or ""),
            j.get("title") or j.get("url") or "",
            j.get("ats") or "",
            j.get("status") or "",
            j.get("url") or "",
            j.get("notes") or "",
            (j.get("first_seen") or "")[:10],
        ])

    for i, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
